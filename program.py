import io
import os
import shutil
import sys

import requests
from PIL import Image
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import (
    OneCellAnchor,
    AnchorMarker,
    XDRPositiveSize2D,
)
from openpyxl.utils.units import pixels_to_EMU, points_to_pixels

# ----------------------------- НАСТРОЙКИ ------------------------------------
FIRST_DATA_ROW = 9          # строка, с которой начинаются товары
URL_COLUMN = "B"            # столбец со ссылками на картинки
IMG_SIZE_PX = 200            # итоговый размер картинки, пикс. (200x200)
ROW_HEIGHT_PT = 225          # высота строк начиная с FIRST_DATA_ROW, пункты
# Ширина столбца B = ширине картинки. Единицы ширины столбца в Excel — это
# "символы", не пиксели, поэтому число пересчитывается из IMG_SIZE_PX по формуле,
# обратной той, что использует сам Excel (пиксели ≈ width*7 + 5).
COL_WIDTH = round((IMG_SIZE_PX - 5) / 7, 2)   # при IMG_SIZE_PX=200 → 27.86 (≈200 px)
REQUEST_TIMEOUT = 15          # сек. ожидания ответа на скачивание одной картинки
JPEG_QUALITY = 92
SUCCESS_PHRASE = "🐍 ПИТОНчик умничка и обработал строку {row}"
# -----------------------------------------------------------------------------

COL_B_INDEX0 = 1  # 0-based индекс столбца B для AnchorMarker (A=0, B=1, ...)


def col_width_to_pixels(width: float) -> int:
    """Приближённый перевод ширины столбца Excel (в 'символах') в пиксели —
    та же формула, что использует сам Excel для дефолтного шрифта Calibri 11."""
    return int(round(width * 7 + 5))


def download_image_bytes(url: str) -> bytes:
    resp = requests.get(url, timeout=REQUEST_TIMEOUT)
    resp.raise_for_status()
    if not resp.content:
        raise ValueError("пустой ответ от сервера")
    return resp.content


def to_jpeg_square(raw_bytes: bytes, size: int = IMG_SIZE_PX):
    """Открывает картинку из байтов, определяет реальный формат, приводит
    к RGB (разворачивая прозрачность на белый фон) и уменьшает до size x size.
    Возвращает (BytesIO с JPEG, исходный_формат)."""
    img = Image.open(io.BytesIO(raw_bytes))
    img.load()  # заставляет Pillow действительно прочитать/провалидировать файл
    original_format = img.format or "UNKNOWN"

    if img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info):
        img = img.convert("RGBA")
        background = Image.new("RGB", img.size, (255, 255, 255))
        background.paste(img, mask=img.split()[-1])
        img = background
    else:
        img = img.convert("RGB")

    img = img.resize((size, size), Image.LANCZOS)

    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=JPEG_QUALITY)
    buf.seek(0)
    return buf, original_format


def make_centered_anchor(row_idx0: int, cell_w_px: int, cell_h_px: int, size: int = IMG_SIZE_PX):
    off_x = max(0, (cell_w_px - size) // 2)
    off_y = max(0, (cell_h_px - size) // 2)
    marker = AnchorMarker(
        col=COL_B_INDEX0, colOff=pixels_to_EMU(off_x),
        row=row_idx0, rowOff=pixels_to_EMU(off_y),
    )
    ext = XDRPositiveSize2D(cx=pixels_to_EMU(size), cy=pixels_to_EMU(size))
    return OneCellAnchor(_from=marker, ext=ext)


def process(source_file: str, output_file: str) -> int:
    if not os.path.exists(source_file):
        print(f"Не найден исходный файл: {source_file}")
        return 1

    shutil.copyfile(source_file, output_file)  # копия сохраняет 100% форматирования

    wb = load_workbook(output_file)
    ws = wb.active

    ws.column_dimensions[URL_COLUMN].width = COL_WIDTH
    cell_w_px = col_width_to_pixels(COL_WIDTH)
    cell_h_px = points_to_pixels(ROW_HEIGHT_PT)

    processed = 0
    skipped = 0
    error_rows = []

    max_row = ws.max_row
    row = FIRST_DATA_ROW
    while row <= max_row:
        # строка считается "реальной", если хоть что-то есть в A..G
        row_values = [ws.cell(row=row, column=c).value for c in range(1, 8)]
        has_any_value = any(v is not None and str(v).strip() != "" for v in row_values)
        if not has_any_value:
            row += 1
            continue

        # высота строки выставляется ВСЕМ строкам с данными, независимо от наличия фото
        ws.row_dimensions[row].height = ROW_HEIGHT_PT

        b_cell = ws[f"{URL_COLUMN}{row}"]
        b_val = b_cell.value
        if b_val is None or str(b_val).strip() == "":
            skipped += 1
            row += 1
            continue

        first_url = str(b_val).split(",")[0].strip()

        try:
            raw = download_image_bytes(first_url)
            jpeg_buf, _orig_format = to_jpeg_square(raw, IMG_SIZE_PX)

            xl_img = XLImage(jpeg_buf)
            xl_img.width = IMG_SIZE_PX
            xl_img.height = IMG_SIZE_PX
            xl_img.anchor = make_centered_anchor(row - 1, cell_w_px, cell_h_px, IMG_SIZE_PX)

            b_cell.value = None  # удаляем текст ячейки — вместо него теперь картинка
            ws.add_image(xl_img)

            processed += 1
            print(SUCCESS_PHRASE.format(row=row))
        except Exception as exc:
            error_rows.append(row)
            print(f"ERROR: строка {row} — не удалось получить/обработать картинку по ссылке '{first_url}' ({exc})")

        row += 1

    wb.save(output_file)

    print("\n----------------- СТАТИСТИКА -----------------")
    print(f"Успешно обработано строк (картинка вставлена): {processed}")
    print(f"Пропущено строк без ссылки на фото:            {skipped}")
    print(f"Ошибок скачивания/обработки:                    {len(error_rows)}")
    if error_rows:
        print(f"  Строки с ошибками: {', '.join(map(str, error_rows))}")
    print(f"Файл сохранён как: {output_file}")
    print("-----------------------------------------------")
    return 0


def main():
    source_file = sys.argv[1] if len(sys.argv) > 1 else "prom.xlsx"
    output_file = sys.argv[2] if len(sys.argv) > 2 else "EUROMAX-PRICE_.xlsx"
    code = process(source_file, output_file)
    try:
        input("\nНажмите Enter, чтобы закрыть окно...")
    except EOFError:
        pass
    sys.exit(code)


if __name__ == "__main__":
    main()
